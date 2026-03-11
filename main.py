from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from pydantic import BaseModel
from groq import Groq
import os, io, base64, tempfile

# مكتبات الملفات
import pypdf
import pdfplumber
from docx import Document
import openpyxl
import pandas as pd
import numpy as np
from PIL import Image

# مكتبات جديدة
from duckduckgo_search import DDGS
from youtube_transcript_api import YouTubeTranscriptApi
from gtts import gTTS

try:
    import pytesseract
    OCR_AVAILABLE = True
except:
    OCR_AVAILABLE = False

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

client = Groq(api_key=os.environ.get("GROQ_API_KEY"))

# ══ استخراج النص من الملفات ══

def extract_pdf(data: bytes) -> str:
    text = ""
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
                tables = page.extract_tables()
                for table in tables:
                    for row in table:
                        text += " | ".join([str(c) for c in row if c]) + "\n"
    except:
        reader = pypdf.PdfReader(io.BytesIO(data))
        for page in reader.pages:
            text += page.extract_text() or ""
    return text[:8000]

def extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    for table in doc.tables:
        for row in table.rows:
            text += " | ".join([c.text.strip() for c in row.cells]) + "\n"
    return text[:8000]

def extract_excel(data: bytes) -> str:
    try:
        df_dict = pd.read_excel(io.BytesIO(data), sheet_name=None)
        text = ""
        for sheet_name, df in df_dict.items():
            text += f"\n[ورقة: {sheet_name}]\n"
            text += f"الأبعاد: {df.shape[0]} صف × {df.shape[1]} عمود\n"
            text += f"الأعمدة: {', '.join([str(c) for c in df.columns])}\n"
            text += df.to_string(max_rows=50) + "\n"
            # إحصائيات
            numeric_cols = df.select_dtypes(include=[np.number])
            if not numeric_cols.empty:
                text += "\nإحصائيات:\n" + numeric_cols.describe().to_string() + "\n"
        return text[:8000]
    except:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        text = ""
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            text += f"\n[ورقة: {sheet}]\n"
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None]
                if vals:
                    text += " | ".join(vals) + "\n"
        return text[:8000]

def extract_csv(data: bytes) -> str:
    try:
        df = pd.read_csv(io.BytesIO(data))
        text = f"عدد الصفوف: {len(df)}\n"
        text += f"الأعمدة: {', '.join(df.columns)}\n"
        text += df.to_string(max_rows=50) + "\n"
        numeric_cols = df.select_dtypes(include=[np.number])
        if not numeric_cols.empty:
            text += "\nإحصائيات:\n" + numeric_cols.describe().to_string()
        return text[:8000]
    except:
        return data.decode("utf-8", errors="ignore")[:8000]

def extract_image(data: bytes, filename: str) -> str:
    # OCR أولاً إذا كان متاحاً
    if OCR_AVAILABLE:
        try:
            img = Image.open(io.BytesIO(data))
            text = pytesseract.image_to_string(img, lang='ara+eng')
            if text.strip():
                return f"[نص مستخرج من الصورة عبر OCR]\n{text}"
        except:
            pass
    return "__IMAGE__"

def extract_text(filename: str, data: bytes) -> str:
    ext = filename.lower().split(".")[-1]
    if ext == "pdf":                    return extract_pdf(data)
    elif ext == "docx":                 return extract_docx(data)
    elif ext in ("xlsx", "xls"):        return extract_excel(data)
    elif ext == "csv":                  return extract_csv(data)
    elif ext in ("jpg","jpeg","png","webp"): return extract_image(data, filename)
    elif ext in ("txt", "md"):
        return data.decode("utf-8", errors="ignore")[:8000]
    return ""

# ══ نقاط API ══

class ChatRequest(BaseModel):
    messages: list
    model: str = "llama-3.3-70b-versatile"

class SearchRequest(BaseModel):
    query: str
    max_results: int = 5

class TTSRequest(BaseModel):
    text: str
    lang: str = "ar"

class YouTubeRequest(BaseModel):
    url: str

@app.post("/api/chat")
async def chat(req: ChatRequest):
    response = client.chat.completions.create(
        model=req.model,
        messages=req.messages,
        max_tokens=1024,
    )
    return {"reply": response.choices[0].message.content}

@app.post("/api/upload")
async def upload(file: UploadFile = File(...)):
    data = await file.read()
    text = extract_text(file.filename, data)

    if text == "__IMAGE__":
        b64 = base64.b64encode(data).decode()
        ext = file.filename.lower().split(".")[-1]
        mime = "image/jpeg" if ext in ("jpg","jpeg") else f"image/{ext}"
        response = client.chat.completions.create(
            model="llama-3.2-11b-vision-preview",
            messages=[{
                "role": "user",
                "content": [
                    {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
                    {"type": "text", "text": "صف محتوى هذه الصورة بالتفصيل بالعربية"}
                ]
            }],
            max_tokens=1024
        )
        return {"text": response.choices[0].message.content, "filename": file.filename, "type": "image"}

    if not text:
        return {"text": "", "filename": file.filename, "message": "لم يتم استخراج نص"}

    return {"text": text, "filename": file.filename, "type": "document"}

@app.post("/api/search")
async def search(req: SearchRequest):
    try:
        results = []
        with DDGS() as ddgs:
            for r in ddgs.text(req.query, max_results=req.max_results):
                results.append(f"• {r['title']}\n{r['body']}\n{r['href']}")
        text = "\n\n".join(results)
        return {"results": text, "count": len(results)}
    except Exception as e:
        return {"results": "", "error": str(e)}

@app.post("/api/youtube")
async def youtube(req: YouTubeRequest):
    try:
        # استخراج معرف الفيديو
        url = req.url
        video_id = ""
        if "youtu.be/" in url:
            video_id = url.split("youtu.be/")[1].split("?")[0]
        elif "v=" in url:
            video_id = url.split("v=")[1].split("&")[0]

        if not video_id:
            return {"transcript": "", "error": "رابط غير صحيح"}

        transcript_list = YouTubeTranscriptApi.get_transcript(
            video_id, languages=["ar", "en", "auto"]
        )
        transcript = " ".join([t["text"] for t in transcript_list])
        return {"transcript": transcript[:8000], "video_id": video_id}
    except Exception as e:
        return {"transcript": "", "error": str(e)}

@app.post("/api/tts")
async def tts(req: TTSRequest):
    try:
        tts_obj = gTTS(text=req.text[:500], lang=req.lang, slow=False)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".mp3")
        tts_obj.save(tmp.name)
        return FileResponse(tmp.name, media_type="audio/mpeg", filename="reply.mp3")
    except Exception as e:
        return {"error": str(e)}

@app.get("/health")
def health():
    return {"status": "ok", "ocr": OCR_AVAILABLE}

app.mount("/", StaticFiles(directory="/app/frontend", html=True), name="frontend")
